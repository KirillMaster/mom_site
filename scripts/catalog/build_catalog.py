import json, os, re, collections, urllib.request
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

SITE = "https://angelamoiseenko.ru"
OUT = os.environ.get("CATALOG_OUT", "catalog.xlsx")

data = json.load(urllib.request.urlopen(f"{SITE}/api/public/gallery", timeout=60))
arts = data["artworks"]

TR = dict(zip("абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
              ["a","b","v","g","d","e","e","zh","z","i","y","k","l","m","n","o","p","r","s","t","u","f","h","ts","ch","sh","sch","","y","","e","yu","ya"]))
def slug(title, id_):
    s = "".join(TR.get(c, c) for c in (title or "").lower())
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-") or "artwork"
    return f"{s}-{id_}"

TECH_WORDS = r"масло|акварель|пастель|темпера|гуашь|акрил|смешанная техника|графика|уголь|карандаш"
SIZE_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*[хxX×]\s*(\d+(?:[.,]\d+)?)")
YEAR_RE = re.compile(r"\b(19[5-9]\d|20[0-2]\d)\s*(?:г|год)?")

def num(s):
    v = float(s.replace(",", "."))
    return int(v) if v.is_integer() else v

def parse(desc):
    desc = (desc or "").replace("\r", "")
    lines = [l.strip() for l in desc.split("\n") if l.strip()]
    meta, rest = [], []
    for i, l in enumerate(lines):
        if len(l) < 90 and not rest:
            meta.append(l)
        else:
            if not rest and len(l) >= 90 and SIZE_RE.search(l[:80]):
                head, _, tail = l.partition(".")
                m = re.match(r"^(.*?\d{4}\s*(?:г\.?|год\.?)?)\s*(.*)$", l)
                if m and len(m.group(1)) < 90:
                    meta.append(m.group(1)); rest.append(m.group(2)); continue
            rest.append(l)
    metatext = ", ".join(meta)
    base = ""
    first = re.split(r",", metatext)[0].strip() if metatext else ""
    if re.search(r"холст|картон|бумага|оргалит|дерев|ДВП", first, re.I):
        base = first
    tech = ", ".join(dict.fromkeys(m.group(0) for m in re.finditer(TECH_WORDS, metatext, re.I)))
    w = h = ""
    m = SIZE_RE.search(metatext)
    if m: w, h = num(m.group(1)), num(m.group(2))
    y = ""
    m = YEAR_RE.search(metatext)
    if m: y = int(m.group(1))
    status = "В наличии" if re.search(r"в наличии", desc, re.I) else ""
    return base, tech, w, h, y, status, "\n".join(rest).strip(), desc.strip()

STATUSES = ["В наличии", "Продана", "Частная коллекция", "Недоступна", "Не продаётся", "Не моя работа"]

thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="3B4A6B")
EDIT_FILL = PatternFill("solid", fgColor="FFF8E1")
AUTO_FILL = PatternFill("solid", fgColor="EEF4FB")
RO_FILL = PatternFill("solid", fgColor="F2F2F2")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

# --- Инструкция
ins = wb.active
ins.title = "Как заполнять"
guide = [
    "Каталог картин — как заполнять",
    "",
    "Одна строка — одна картина. Все работы с сайта уже здесь.",
    "",
    "🟨 Жёлтые колонки — заполняешь ты: цена, статус, описание, история, ⭐.",
    "🟦 Голубые — я заполнил сам из текста на сайте (размер, техника, год). Проверь и поправь, если не так.",
    "⬜ Серые — справочно, не трогай (ID, ссылка, текущий текст на сайте).",
    "",
    "Цена — в рублях, только цифры. Если не хочешь публиковать — оставь пусто и напиши «по запросу» в комментарии.",
    "Размер — ширина × высота в см. Если на сайте было наоборот — поменяй местами.",
    "Статус — выбери из списка: В наличии / Продана / Частная коллекция / Недоступна / Не продаётся / Не моя работа.",
    "⭐ — поставь «да» напротив самых сильных работ.",
    "«Возможный дубль» — на сайте несколько работ с таким же названием. Если это одна и та же картина — напиши в комментарии «дубль».",
    "Фото — если на фото блики или неправильный цвет, отметь «переснять».",
    "",
    "Если чего-то не знаешь — оставляй пусто, это нормально.",
    "Новые работы, которых нет на сайте, добавляй в конец таблицы (ID оставь пустым).",
    "",
    "Отдельный лист «Фото с выставок» — это не картины, а фото. Отметь, какие оставить на сайте.",
    "Лист «О себе» — биография, выставки, награды.",
    "",
    "format: v1 (не удаляй эту строку)",
]
for i, t in enumerate(guide, 1):
    c = ins.cell(row=i, column=1, value=t)
    c.alignment = Alignment(wrap_text=True)
ins["A1"].font = Font(bold=True, size=16)
ins.column_dimensions["A"].width = 110

# --- Каталог
ws = wb.create_sheet("Каталог")
cols = [
    ("ID", 7, RO_FILL), ("Раздел на сайте", 16, RO_FILL), ("Фото", 16, RO_FILL), ("Название", 26, EDIT_FILL),
    ("Цена, ₽", 12, EDIT_FILL), ("Статус", 17, EDIT_FILL),
    ("Ширина, см", 10, AUTO_FILL), ("Высота, см", 10, AUTO_FILL), ("Год", 8, AUTO_FILL),
    ("Основа", 18, AUTO_FILL), ("Техника", 16, AUTO_FILL),
    ("Короткое описание (1–2 предложения)", 40, EDIT_FILL),
    ("История, выставки, интересные факты", 40, EDIT_FILL),
    ("⭐ Сильная работа", 10, EDIT_FILL), ("Фото: переснять?", 11, EDIT_FILL),
    ("Комментарий", 30, EDIT_FILL), ("Возможный дубль (ID)", 14, RO_FILL),
    ("Ссылка на сайте", 30, RO_FILL), ("Текущий текст на сайте", 50, RO_FILL),
]
for j, (name, width, fill) in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=name)
    c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.alignment = center; c.border = border
    ws.column_dimensions[c.column_letter].width = width
ws.row_dimensions[1].height = 45
ws.freeze_panes = "E2"

paint = [a for a in arts if a["category"]["name"] != "Фото с выставок"]
expo = [a for a in arts if a["category"]["name"] == "Фото с выставок"]
paint.sort(key=lambda a: (a["category"].get("displayOrder", 0), a["title"].strip('"«» ').lower(), a["id"]))

by_title = collections.defaultdict(list)
for a in paint:
    by_title[a["title"].strip('"«» ').lower()].append(a["id"])

stats = collections.Counter()
for i, a in enumerate(paint, 2):
    base, tech, w, h, y, status, rest, full = parse(a.get("description"))
    stats["size"] += bool(w); stats["year"] += bool(y); stats["tech"] += bool(tech); stats["status"] += bool(status)
    dups = [x for x in by_title[a["title"].strip('"«» ').lower()] if x != a["id"]]
    url = f"{SITE}/gallery/{slug(a['title'], a['id'])}"
    row = [a["id"], a["category"]["name"], f'=IMAGE("{a["thumbnailPath"]}")', a["title"].strip(),
           a.get("price") or None, status or None, w or None, h or None, y or None, base or None, tech or None,
           None, None, None, None, None, ", ".join(map(str, dups)) or None, url, full or None]
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.fill = cols[j-1][2]; c.border = border; c.alignment = wrap
    ws.cell(row=i, column=18).hyperlink = url
    ws.cell(row=i, column=5).number_format = '#,##0 "₽"'
    ws.row_dimensions[i].height = 90

last = len(paint) + 1
extra = last + 200
for j in range(1, len(cols) + 1):
    for i in range(last + 1, extra):
        ws.cell(row=i, column=j).fill = cols[j-1][2] if cols[j-1][2] is not RO_FILL else EDIT_FILL
dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
dv_star = DataValidation(type="list", formula1='"да"', allow_blank=True)
dv_photo = DataValidation(type="list", formula1='"переснять"', allow_blank=True)
for d in (dv, dv_star, dv_photo): ws.add_data_validation(d)
dv.add(f"F2:F{extra}"); dv_star.add(f"N2:N{extra}"); dv_photo.add(f"O2:O{extra}")
ws.conditional_formatting.add(f"A2:S{extra}", FormulaRule(formula=['$F2="Продана"'], fill=PatternFill("solid", fgColor="E0E0E0")))
ws.conditional_formatting.add(f"A2:S{extra}", FormulaRule(formula=['$N2="да"'], fill=PatternFill("solid", fgColor="FFE9A8")))
ws.auto_filter.ref = f"A1:S{last}"

# --- Фото с выставок
ex = wb.create_sheet("Фото с выставок")
ecols = [("ID", 7), ("Фото", 16), ("Подпись на сайте", 40), ("Оставить на сайте?", 14), ("Где и когда (выставка, год)", 40), ("Ссылка", 30)]
for j, (n, w) in enumerate(ecols, 1):
    c = ex.cell(row=1, column=j, value=n); c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.alignment = center
    ex.column_dimensions[c.column_letter].width = w
ex.freeze_panes = "C2"
dv_keep = DataValidation(type="list", formula1='"да,нет"', allow_blank=True); ex.add_data_validation(dv_keep)
for i, a in enumerate(sorted(expo, key=lambda a: a["id"]), 2):
    url = f"{SITE}/gallery/{slug(a['title'], a['id'])}"
    for j, v in enumerate([a["id"], f'=IMAGE("{a["thumbnailPath"]}")', a["title"].strip(), None, (a.get("description") or "").strip() or None, url], 1):
        c = ex.cell(row=i, column=j, value=v); c.alignment = wrap; c.border = border
        c.fill = EDIT_FILL if j in (4, 5) else RO_FILL
    ex.row_dimensions[i].height = 90
dv_keep.add(f"D2:D{len(expo)+1}")

# --- О себе
ab = wb.create_sheet("О себе")
for j, (n, w) in enumerate([("Раздел", 28), ("Твой текст", 100)], 1):
    c = ab.cell(row=1, column=j, value=n); c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD_FILL
    ab.column_dimensions[c.column_letter].width = w
for i, n in enumerate(["Биография (своими словами)", "Образование", "Выставки (год, место, название)",
                       "Награды, звания, членство", "Публикации, СМИ, интервью", "Где находятся работы (музеи, коллекции)",
                       "Мастер-классы: опыт", "Отзывы покупателей/учеников (или ссылки на скрины)"], 2):
    ab.cell(row=i, column=1, value=n).font = Font(bold=True)
    c = ab.cell(row=i, column=2); c.fill = EDIT_FILL; c.alignment = wrap
    ab.row_dimensions[i].height = 120

wb.active = 1
wb.save(OUT)
print(OUT, "paintings:", len(paint), "expo:", len(expo), dict(stats),
      "dup rows:", sum(1 for v in by_title.values() if len(v) > 1 for _ in v))
